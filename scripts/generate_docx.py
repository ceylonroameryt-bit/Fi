import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(WORKSPACE_DIR, "docs")
STANDARD_DOCS_DIR = os.path.join(DOCS_DIR, "15-standard-docs")
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")
OUTPUT_DOCX = os.path.join(DOCS_DIR, "blynt_Complete_Documentation.docx")
LOGO_PATH = os.path.join(WORKSPACE_DIR, "blynt-logo-horizontal.png")

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(15)

    if os.path.exists(LOGO_PATH):
        try:
            run = p.add_run()
            run.add_picture(LOGO_PATH, width=Inches(3.2))
        except Exception:
            pass

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(25)
    p_title.paragraph_format.space_after = Pt(10)
    run_title = p_title.add_run("blynt")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(36)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(22, 45, 95) # Deep UK Navy #162D5F

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Comprehensive Master Specification, Business Blueprint\n& Technical Architecture Explainer")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(17)
    run_sub.font.color.rgb = RGBColor(51, 65, 85)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(35)
    r_line = p_line.add_run("— Executive Non-Technical Guide & Full Engineering Blueprints with High-Resolution Architecture Diagrams —")
    r_line.font.name = "Segoe UI"
    r_line.font.size = Pt(10.5)
    r_line.font.italic = True
    r_line.font.color.rgb = RGBColor(255, 127, 108) # Coral #FF7F6C

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(45)
    
    meta_lines = [
        ("Product / Platform", "blynt (Internal Code: Ledgerline)"),
        ("Source Code Repository", "https://github.com/ceylonroameryt-bit/Fi.git"),
        ("Target Jurisdiction", "United Kingdom (HMRC • Companies House • UK GDPR)"),
        ("Technology Stack", "NestJS 11 • PostgreSQL 16 • Prisma ORM • Next.js 15"),
        ("Standard Compliance", "UK GAAP • Companies Act 2006 • Making Tax Digital (MTD)"),
        ("Document Version", "2.0 (Complete Executive & Technical Publication)"),
        ("Published Date", "October 2026"),
        ("Authors", "Founding Team & Antigravity AI Engineering"),
    ]
    for label, val in meta_lines:
        r_lbl = p_meta.add_run(f"{label}: ")
        r_lbl.font.name = "Segoe UI"
        r_lbl.font.size = Pt(10)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(22, 45, 95)
        
        r_v = p_meta.add_run(f"{val}\n")
        r_v.font.name = "Segoe UI"
        r_v.font.size = Pt(10)
        r_v.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

def format_inline_text(paragraph, text):
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
            run.font.name = "Segoe UI"
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
            run.font.name = "Segoe UI"
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(190, 24, 93) # Deep Pink/Red code
        else:
            run = paragraph.add_run(token)
            run.font.name = "Segoe UI"

def resolve_image_path(raw_path):
    # Try direct
    if os.path.exists(raw_path):
        return raw_path
    # Try relative to diagrams dir
    p1 = os.path.join(DIAGRAMS_DIR, os.path.basename(raw_path))
    if os.path.exists(p1):
        return p1
    # Try relative to docs dir
    p2 = os.path.join(DOCS_DIR, raw_path)
    if os.path.exists(p2):
        return p2
    # Try relative to workspace
    p3 = os.path.join(WORKSPACE_DIR, raw_path)
    if os.path.exists(p3):
        return p3
    return None

def parse_markdown_to_docx(doc, md_content):
    lines = md_content.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        num_cols = max(len(r) for r in table_rows)
        if num_cols == 0:
            in_table = False
            table_rows = []
            return

        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row_data in enumerate(table_rows):
            is_header = (r_idx == 0)
            row = table.rows[r_idx]
            for c_idx in range(num_cols):
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell = row.cells[c_idx]
                set_cell_margins(cell, top=70, bottom=70, left=110, right=110)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                if is_header:
                    set_cell_background(cell, "162D5F") # Navy header
                    run = p.add_run(cell_text.strip())
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(9.5)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    format_inline_text(p, cell_text.strip())
                    for run in p.runs:
                        run.font.size = Pt(9)
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)
        in_table = False
        table_rows = []

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)
        in_code_block = False
        code_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle image ![caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match and not in_code_block:
            if in_table:
                flush_table()
            caption, img_raw = img_match.groups()
            resolved = resolve_image_path(img_raw)
            if resolved:
                try:
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(10)
                    p_img.paragraph_format.space_after = Pt(2)
                    r_img = p_img.add_run()
                    r_img.add_picture(resolved, width=Inches(6.2))
                    
                    if caption:
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(2)
                        p_cap.paragraph_format.space_after = Pt(12)
                        r_cap = p_cap.add_run(f"Figure: {caption}")
                        r_cap.font.name = "Segoe UI"
                        r_cap.font.size = Pt(9.5)
                        r_cap.font.italic = True
                        r_cap.font.color.rgb = RGBColor(100, 116, 139)
                except Exception as e:
                    print(f"Error embedding image {resolved}: {e}")
            continue

        # Code block handler
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Table handler
        if stripped.startswith("|") and stripped.endswith("|"):
            parts = [p.strip() for p in stripped.split("|")[1:-1]]
            if all(set(p).issubset({'-', ':', ' '}) and len(p) > 0 for p in parts):
                continue
            in_table = True
            table_rows.append(parts)
            continue
        elif in_table:
            flush_table()

        if not stripped:
            continue

        # Horizontal Rule
        if stripped in ["---", "***", "___"]:
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
            continue

        # Blockquote / Callout box
        if stripped.startswith("> "):
            quote_text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "EEF2F9")
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, quote_text)
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.italic = True
                r.font.color.rgb = RGBColor(22, 45, 95)
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(4)
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(20)
            h.paragraph_format.space_after = Pt(8)
            run = h.add_run(stripped[2:])
            run.font.name = "Segoe UI"
            run.font.bold = True
            run.font.size = Pt(21)
            run.font.color.rgb = RGBColor(22, 45, 95) # Deep Navy
            continue
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[3:])
            run.font.name = "Segoe UI"
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(255, 127, 108) # Brand Coral Accent #FF7F6C
            continue
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[4:])
            run.font.name = "Segoe UI"
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(71, 85, 105) # Slate
            continue

        # Bullet list
        if stripped.startswith("* ") or stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, stripped[2:])
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, num_match.group(2))
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        format_inline_text(p, stripped)

    if in_code_block:
        flush_code()
    if in_table:
        flush_table()

def main():
    doc = Document()
    
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    print("Adding cover page...")
    add_cover_page(doc)

    doc_files = [
        ("01_PROJECT_CHARTER.md", "1. Project Charter"),
        ("02_BUSINESS_REQUIREMENTS_DOCUMENT.md", "2. Business Requirements Document (BRD)"),
        ("03_PRODUCT_REQUIREMENTS_DOCUMENT.md", "3. Product Requirements Document (PRD)"),
        ("04_USER_STORIES_AND_ACCEPTANCE_CRITERIA.md", "4. User Stories & Acceptance Criteria"),
        ("05_SYSTEM_ARCHITECTURE_DOCUMENT.md", "5. System Architecture Document (SAD / TDD)"),
        ("06_DATABASE_SCHEMA_AND_ERD.md", "6. Database Schema & ERD"),
        ("07_API_SPECIFICATION.md", "7. REST API Specification"),
        ("08_ARCHITECTURE_DECISION_RECORDS.md", "8. Architecture Decision Records (ADRs)"),
        ("09_UI_UX_DESIGN_SPECIFICATION.md", "9. UI/UX Design System Specification"),
        ("10_THREAT_MODEL_AND_SECURITY_PLAN.md", "10. Threat Model & Security Plan"),
        ("11_COMPLIANCE_AND_PRIVACY_CHECKLIST.md", "11. Compliance & Privacy Checklist"),
        ("12_DEPLOYMENT_AND_DEVOPS_RUNBOOK.md", "12. Deployment & DevOps Runbook"),
        ("13_PROJECT_ROADMAP_AND_RELEASE_PLAN.md", "13. Project Roadmap & Release Plan"),
        ("14_RISK_REGISTER_RAID_LOG.md", "14. Risk Register & RAID Log"),
        ("15_RACI_MATRIX.md", "15. RACI Governance Matrix"),
    ]

    for fname, title in doc_files:
        fpath = os.path.join(STANDARD_DOCS_DIR, fname)
        if os.path.exists(fpath):
            print(f"Processing {fname}...")
            with open(fpath, "r", encoding="utf-8") as f:
                content = f.read()
            parse_markdown_to_docx(doc, content)
            doc.add_page_break()

    tech_explainer = os.path.join(DOCS_DIR, "TECHNICAL_ARCHITECTURE_AND_CODE_EXPLAINER.md")
    if os.path.exists(tech_explainer):
        print("Processing TECHNICAL_ARCHITECTURE_AND_CODE_EXPLAINER.md...")
        with open(tech_explainer, "r", encoding="utf-8") as f:
            content = f.read()
        parse_markdown_to_docx(doc, content)

    print(f"Saving to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print("Document successfully created!")

if __name__ == "__main__":
    main()
