import os
import re
import sys
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(WORKSPACE_DIR, "docs")
SOURCE_DIR = os.path.join(DOCS_DIR, "15-standard-docs")
OUTPUT_DOCX = os.path.join(DOCS_DIR, "Blynt_Complete_15_Standard_Project_Documents.docx")
OUTPUT_PDF = os.path.join(DOCS_DIR, "Blynt_Complete_15_Standard_Project_Documents.pdf")
LOGO_PATH = os.path.join(DOCS_DIR, "blynt-logo.png")
if not os.path.exists(LOGO_PATH):
    LOGO_PATH = os.path.join(WORKSPACE_DIR, "blynt-logo.png")

NAVY = RGBColor(22, 45, 95)       # #162D5F
CORAL = RGBColor(255, 127, 108)   # #FF7F6C
SLATE = RGBColor(71, 85, 105)     # #475569
CHARCOAL = RGBColor(30, 41, 59)   # #1E293B

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(10)

    if os.path.exists(LOGO_PATH):
        try:
            run = p.add_run()
            run.add_picture(LOGO_PATH, width=Inches(1.8))
        except Exception:
            pass

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(25)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("blynt")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(40)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("COMPLETE 15-STANDARD PROJECT MANAGEMENT & GOVERNANCE SUITE")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(16)
    run_sub.font.bold = True
    run_sub.font.color.rgb = NAVY

    p_tagline = doc.add_paragraph()
    p_tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tagline.paragraph_format.space_after = Pt(25)
    r_tagline = p_tagline.add_run("Project Charter • Business & Product Requirements • Architecture & Database\nThreat Model & Security • DevOps & Runbooks • Roadmap, RAID & RACI Governance\nOfficial Engineering Delivery Compendium • October 2026")
    r_tagline.font.name = "Segoe UI"
    r_tagline.font.size = Pt(10.5)
    r_tagline.font.color.rgb = CORAL

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_background(c, "EEF2F9")
    set_cell_margins(c, top=120, bottom=120, left=200, right=200)
    p_in = c.paragraphs[0]
    p_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_in = p_in.add_run("Executive PMO & Architectural Compendium\n")
    r_in.font.name = "Segoe UI"
    r_in.font.size = Pt(12)
    r_in.font.bold = True
    r_in.font.color.rgb = NAVY
    r_in2 = p_in.add_run("This master documentation suite compiles all 15 industry-standard project management, architectural, and governance deliverables for the blynt enterprise double-entry accounting platform. It serves as the authoritative single source of truth (SSOT) across executive, product, engineering, security, and compliance directorates.")
    r_in2.font.name = "Segoe UI"
    r_in2.font.size = Pt(9.5)
    r_in2.font.italic = True
    r_in2.font.color.rgb = SLATE

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(35)
    
    meta_lines = [
        ("Author & Architect", "Principal Systems Architect & PMO Directorate"),
        ("Ecosystem", "NestJS 11 • Next.js 15 • PostgreSQL 16 • AWS App Runner (eu-west-2)"),
        ("Compliance Scope", "UK GAAP • HMRC MTD • SOC 2 Type II • ISO 27001 • GDPR"),
        ("Document Reference", "BLYNT-PMO-15-STANDARDS-V1"),
        ("Publication Date", "October 2026 • Production Verified")
    ]
    for lbl, val in meta_lines:
        r1 = p_meta.add_run(f"{lbl}: ")
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r2 = p_meta.add_run(f"{val}\n")
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = SLATE

    doc.add_page_break()

def parse_markdown_to_docx(doc, md_content, doc_num, doc_title):
    # Section Header
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(20)
    p_sec.paragraph_format.space_after = Pt(10)
    r_sec = p_sec.add_run(f"PART {doc_num:02d} — {doc_title.upper()}")
    r_sec.font.name = "Segoe UI"
    r_sec.font.size = Pt(18)
    r_sec.font.bold = True
    r_sec.font.color.rgb = NAVY

    lines = md_content.splitlines()
    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()

        # Table detection
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # End of table, process it
            render_markdown_table(doc, table_lines)
            in_table = False
            table_lines = []

        if not stripped:
            continue

        # Skip main markdown H1 since we added section header
        if stripped.startswith("# "):
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(stripped[3:].strip())
            r.font.name = "Segoe UI"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = NAVY

        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(stripped[4:].strip())
            r.font.name = "Segoe UI"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = CORAL

        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(stripped[5:].strip())
            r.font.name = "Segoe UI"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = CHARCOAL

        elif stripped.startswith("* ") or stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, stripped[2:].strip())

        elif re.match(r"^\d+\.\s+", stripped):
            num_content = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, num_content)

        elif stripped.startswith("> "):
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tbl.cell(0, 0)
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, top=60, bottom=60, left=120, right=120)
            p_call = c.paragraphs[0]
            format_inline_text(p_call, stripped[2:].strip(), italic=True)

        elif stripped.startswith("---"):
            continue

        elif stripped.startswith("!["):
            # Image mention - ignore markdown tag or render placeholder
            continue

        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            format_inline_text(p, stripped)

    if in_table and table_lines:
        render_markdown_table(doc, table_lines)

    doc.add_page_break()

def format_inline_text(paragraph, text, italic=False):
    # Simple regex parser for **bold** and `code`
    parts = re.split(r"(\*\*.*?\*\*|\`.*?\`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(10)
            run.font.color.rgb = CHARCOAL
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = NAVY
        else:
            run = paragraph.add_run(part)
            run.font.name = "Segoe UI"
            run.font.size = Pt(10)
            run.font.color.rgb = CHARCOAL
        if italic:
            run.font.italic = True

def render_markdown_table(doc, lines):
    rows_data = []
    for l in lines:
        if re.match(r"^\|?\s*[-:]+[-| :]*$", l):
            continue
        cells = [c.strip() for c in l.strip("|").split("|")]
        if cells:
            rows_data.append(cells)

    if not rows_data:
        return

    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, r_data in enumerate(rows_data):
        for col_idx in range(num_cols):
            val = r_data[col_idx] if col_idx < len(r_data) else ""
            cell = table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            if row_idx == 0:
                set_cell_background(cell, "162D5F")
                p = cell.paragraphs[0]
                run = p.add_run(val)
                run.font.bold = True
                run.font.name = "Segoe UI"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                p = cell.paragraphs[0]
                run = p.add_run(val)
                run.font.name = "Segoe UI"
                run.font.size = Pt(8.5)
                run.font.color.rgb = CHARCOAL

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(4)

def build_compendium():
    print("[*] Initializing Word Document...")
    doc = Document()

    # Set 0.75 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    add_cover_page(doc)

    doc_files = [
        ("01_PROJECT_CHARTER.md", "Project Charter"),
        ("02_BUSINESS_REQUIREMENTS_DOCUMENT.md", "Business Requirements Document (BRD)"),
        ("03_PRODUCT_REQUIREMENTS_DOCUMENT.md", "Product Requirements Document (PRD)"),
        ("04_USER_STORIES_AND_ACCEPTANCE_CRITERIA.md", "User Stories & Acceptance Criteria"),
        ("05_SYSTEM_ARCHITECTURE_DOCUMENT.md", "System Architecture Document (SAD)"),
        ("06_DATABASE_SCHEMA_AND_ERD.md", "Database Schema & Relational ERD"),
        ("07_API_SPECIFICATION.md", "REST & Subledger API Specification"),
        ("08_ARCHITECTURE_DECISION_RECORDS.md", "Architecture Decision Records (ADRs)"),
        ("09_UI_UX_DESIGN_SPECIFICATION.md", "UI/UX Design & Human Interface Spec"),
        ("10_THREAT_MODEL_AND_SECURITY_PLAN.md", "Threat Model & Security Plan"),
        ("11_COMPLIANCE_AND_PRIVACY_CHECKLIST.md", "Compliance, HMRC & Privacy Checklist"),
        ("12_DEPLOYMENT_AND_DEVOPS_RUNBOOK.md", "Deployment & DevOps SRE Runbook"),
        ("13_PROJECT_ROADMAP_AND_RELEASE_PLAN.md", "Project Roadmap & Release Plan"),
        ("14_RISK_REGISTER_RAID_LOG.md", "Risk Register & RAID Governance Log"),
        ("15_RACI_MATRIX.md", "RACI Governance Matrix")
    ]

    for idx, (filename, title) in enumerate(doc_files, 1):
        fpath = os.path.join(SOURCE_DIR, filename)
        if os.path.exists(fpath):
            print(f"[+] Compiling Document {idx}/15: {filename}...")
            with open(fpath, "r", encoding="utf-8") as f:
                content = f.read()
            parse_markdown_to_docx(doc, content, idx, title)
        else:
            print(f"[-] Missing file: {fpath}")

    print(f"[*] Saving master document to: {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print(f"[+] Saved DOCX successfully ({os.path.getsize(OUTPUT_DOCX):,} bytes)")

    # Convert to PDF using Word COM
    print("[*] Converting master document to PDF via Microsoft Word...")
    ps_cmd = f"""
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    try {{
        $doc = $word.Documents.Open('{OUTPUT_DOCX}')
        $doc.SaveAs([ref]'{OUTPUT_PDF}', [ref]17)
        $doc.Close()
        Write-Host "PDF generated: {OUTPUT_PDF}"
    }} finally {{
        $word.Quit()
    }}
    """
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps_cmd], capture_output=True, text=True)
    if os.path.exists(OUTPUT_PDF):
        print(f"[+] Master PDF Generated Successfully! ({os.path.getsize(OUTPUT_PDF):,} bytes)")
        print(f"    -> Location: {OUTPUT_PDF}")
    else:
        print(f"[-] PDF conversion output: {res.stdout} {res.stderr}")

if __name__ == "__main__":
    build_compendium()
